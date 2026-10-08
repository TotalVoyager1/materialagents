import json
from io import BytesIO
from typing import Any

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


def build_markdown(results: dict[str, str], task: str) -> str:
    sections = [
        f"# Итог проекта\n\n**Задача:** {task}",
        ("## План\n\n" + results.get("plan", "")),
        ("## Исследование\n\n" + results.get("research", "")),
        ("## Отчёт\n\n" + results.get("draft", "")),
        ("## Контроль качества\n\n" + results.get("critic", "")),
        ("## Тест-кейсы\n\n" + results.get("tests", "")),
        ("## Проверка кода\n\n" + results.get("code_review", "")),
        ("## Резюме\n\n" + results.get("summary", "")),
    ]
    return "\n\n".join(sections).strip() + "\n"


def build_json(results: dict[str, str], task: str, events: list[dict[str, Any]]) -> bytes:
    return json.dumps(
        {"task": task, "results": results, "events": events},
        ensure_ascii=False,
        indent=2,
    ).encode("utf-8")


def build_docx(markdown: str) -> bytes:
    document = Document()
    for line in markdown.splitlines():
        if line.startswith("# "):
            document.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=1)
        elif line.strip():
            document.add_paragraph(line.replace("**", ""))
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def build_pdf(markdown: str) -> bytes:
    stream = BytesIO()
    document = SimpleDocTemplate(stream, pagesize=A4)
    styles = getSampleStyleSheet()
    story = []
    for line in markdown.splitlines():
        clean = line.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        if not clean.strip():
            story.append(Spacer(1, 6))
        elif clean.startswith("# "):
            story.append(Paragraph(clean[2:], styles["Title"]))
        elif clean.startswith("## "):
            story.append(Paragraph(clean[3:], styles["Heading2"]))
        else:
            story.append(Paragraph(clean.replace("**", ""), styles["BodyText"]))
    document.build(story)
    return stream.getvalue()
