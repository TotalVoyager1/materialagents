from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import re

import numpy as np
import pandas as pd
from pypdf import PdfReader
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
from docx import Document


SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".pdf", ".docx", ".csv", ".py", ".json",
    ".js", ".ts", ".java", ".cs", ".cpp", ".c", ".go", ".rs",
}
MIN_TOKEN_CHAR_SIMILARITY = 0.16


@dataclass
class DocumentChunk:
    filename: str
    text: str


def extract_text(filename: str, data: bytes) -> str:
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Неподдерживаемый формат файла: {extension or '(без расширения)'}")
    if extension == ".pdf":
        reader = PdfReader(BytesIO(data))
        text = "\n".join(page.extract_text() or "" for page in reader.pages)
    elif extension == ".docx":
        document = Document(BytesIO(data))
        paragraphs = [paragraph.text for paragraph in document.paragraphs]
        table_rows = [
            " | ".join(cell.text for cell in row.cells)
            for table in document.tables
            for row in table.rows
        ]
        text = "\n".join(paragraphs + table_rows)
    elif extension == ".csv":
        frame = pd.read_csv(BytesIO(data))
        text = frame.head(100).to_csv(index=False)
    else:
        text = data.decode("utf-8-sig", errors="replace")
    return text.strip()


class DocumentIndex:
    def __init__(self, chunk_size: int = 1200, overlap: int = 150):
        if chunk_size < 1 or overlap < 0 or overlap >= chunk_size:
            raise ValueError("Размер фрагмента должен быть больше перекрытия.")
        self.chunk_size = chunk_size
        self.overlap = overlap
        self.chunks: list[DocumentChunk] = []
        self.word_vectorizer: TfidfVectorizer | None = None
        self.word_matrix = None
        self.char_vectorizer: TfidfVectorizer | None = None
        self.char_matrix = None

    def add(self, filename: str, text: str) -> None:
        clean_text = text.strip()
        if not clean_text:
            return
        step = self.chunk_size - self.overlap
        self.chunks.extend(
            DocumentChunk(filename, clean_text[start : start + self.chunk_size])
            for start in range(0, len(clean_text), step)
        )
        self.word_vectorizer = None
        self.word_matrix = None
        self.char_vectorizer = None
        self.char_matrix = None

    def search(self, query: str, limit: int = 5) -> list[tuple[DocumentChunk, float]]:
        if not self.chunks or not query.strip() or limit < 1:
            return []
        texts = [chunk.text for chunk in self.chunks]
        if self.word_vectorizer is None or self.word_matrix is None:
            self.word_vectorizer = TfidfVectorizer(
                ngram_range=(1, 2),
                sublinear_tf=True,
                token_pattern=r"(?u)\b\w+\b",
            )
            self.word_matrix = self.word_vectorizer.fit_transform(texts)
        if self.char_vectorizer is None or self.char_matrix is None:
            self.char_vectorizer = TfidfVectorizer(
                analyzer="char_wb",
                ngram_range=(2, 5),
                sublinear_tf=True,
            )
            self.char_matrix = self.char_vectorizer.fit_transform(texts)
        word_scores = cosine_similarity(
            self.word_vectorizer.transform([query]), self.word_matrix
        ).ravel()
        char_scores = cosine_similarity(
            self.char_vectorizer.transform([query]), self.char_matrix
        ).ravel()
        query_terms = [
            term for term in re.findall(r"(?u)\b\w+\b", query.lower()) if len(term) > 2
        ]
        unknown_terms = [
            term
            for term in query_terms
            if term not in self.word_vectorizer.vocabulary_
        ]
        if unknown_terms:
            token_char_scores = cosine_similarity(
                self.char_vectorizer.transform(unknown_terms), self.char_matrix
            )
            fuzzy_scores = token_char_scores.max(axis=0)
        else:
            fuzzy_scores = np.zeros(len(texts))
        scores = 0.7 * word_scores + 0.3 * char_scores
        ranked = scores.argsort()[::-1][:limit]
        return [
            (self.chunks[int(i)], float(scores[i]))
            for i in ranked
            if (
                fuzzy_scores[i] >= MIN_TOKEN_CHAR_SIMILARITY
                if unknown_terms
                else word_scores[i] > 0
            )
        ]
